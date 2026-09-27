from pptx import Presentation
from pptx.util import Inches, Pt
import os

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "GradeWise Architecture & Overview"
    subtitle.text = "AI-Powered Jindal Stainless Steel Grade Selection\nReal-time auto-generated presentation"

    # 2. Architecture Slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "System Architecture"
    
    tf = body_shape.text_frame
    tf.text = "A modern decoupled Full-Stack Web Application"
    
    p = tf.add_paragraph()
    p.text = "Frontend: Next.js (React)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Styling & UI: Tailwind CSS, Framer Motion, Lenis, Three.js, Recharts"
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Backend: Python (FastAPI)"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "AI Engine: Groq / Gemini"
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Dataset: JSON-based Stainless Steel database"
    p.level = 1

    # 3. Frontend Details
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Frontend Implementation"
    
    tf = body_shape.text_frame
    tf.text = "Key Components & Features:"
    
    p = tf.add_paragraph()
    p.text = "Framework: Next.js with React Three Fiber (3D graphics)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Typography: Rajdhani, Manrope, Share Tech Mono"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Animations & Experience:"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Morphing Spinners and Preloaders (Framer Motion)"
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Interactive Background Video Component with strict state locking"
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Smooth Scrolling via Lenis (SmoothScrollProvider)"
    p.level = 2

    # 4. Backend & AI Details
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Backend Engine"
    
    tf = body_shape.text_frame
    tf.text = "FastAPI Service Architecture"
    
    p = tf.add_paragraph()
    p.text = "Endpoints: /health, /recommend"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "AI Strategy (Groq/Gemini):"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Reads grades.json to decide on 2-3 grade recommendations"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Strictly natural language trade-off analysis, no hallucinatory numbers"
    p.level = 2
    
    p = tf.add_paragraph()
    p.text = "Physics Engine (calculations.py):"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Provides deterministic numerical outputs (loads, deflections)"
    p.level = 2

    # 5. Dataset Details
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "The Dataset"
    
    tf = body_shape.text_frame
    tf.text = "Jindal Stainless Steel Properties"
    
    p = tf.add_paragraph()
    p.text = "Source: GradeWise_Demo_Steel_Datasheet.xlsx"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Storage: backend/data/grades.json"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Loaded into memory via GRADES_BY_LABEL dictionary"
    p.level = 1

    prs.save('GradeWise_Presentation.pptx')
    print("Presentation saved to GradeWise_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
